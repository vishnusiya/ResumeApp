import io
import os
import re
from dotenv import load_dotenv

# Automatically load environment variables from backend/.env and root .env
load_dotenv()
load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))
load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, JSONResponse
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from .models import (
        ResumeData,
        AIRewriteRequest,
        AIRewriteResponse,
        ATSScoreRequest,
        ResumeMergeRequest,
        ResumeMergeResponse,
        GenerateJDRequest,
        GenerateJDResponse,
    )
    from .services.extractor import extract_text_from_pdf, extract_text_from_docx, parse_raw_text_to_resume
    from .services.ai_agent import run_dual_pass_ai_rewrite, parse_resume_with_ai
    from .services.ats_service import calculate_backend_ats_score
    from .services.merge_engine import run_resume_merger, generate_job_description_from_role
except ImportError:
    from models import (
        ResumeData,
        AIRewriteRequest,
        AIRewriteResponse,
        ATSScoreRequest,
        ResumeMergeRequest,
        ResumeMergeResponse,
        GenerateJDRequest,
        GenerateJDResponse,
    )
    from services.extractor import extract_text_from_pdf, extract_text_from_docx, parse_raw_text_to_resume
    from services.ai_agent import run_dual_pass_ai_rewrite, parse_resume_with_ai
    from services.ats_service import calculate_backend_ats_score
    from services.merge_engine import run_resume_merger, generate_job_description_from_role

app = FastAPI(
    title="AI Resume Editor Backend API",
    version="1.0.0",
    description="Production-ready FastAPI backend with multi-format parsing, Dual-Pass AI agent, and ATS scoring."
)

# CORS middleware for development and production
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/api/health")
async def health_check():
    return {"status": "ok", "service": "AI Resume Editor API", "version": "1.0.0"}

@app.post("/api/upload")
async def upload_resume_file(file: UploadFile = File(...)):
    filename = file.filename or "resume"
    ext = filename.split(".")[-1].lower()
    
    content_bytes = await file.read()

    try:
        if ext == "pdf":
            raw_text = extract_text_from_pdf(content_bytes)
        elif ext in ["docx", "doc"]:
            raw_text = extract_text_from_docx(content_bytes)
        elif ext in ["txt", "md"]:
            raw_text = content_bytes.decode("utf-8", errors="ignore")
        elif ext == "json":
            import json
            data_dict = json.loads(content_bytes.decode("utf-8"))
            return ResumeData(**data_dict)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file format: .{ext}")

        # 1. Try AI-powered LLM parsing if OpenAI API key is available
        if os.environ.get("OPENAI_API_KEY"):
            try:
                parsed_resume = await parse_resume_with_ai(raw_text, filename)
                return parsed_resume
            except Exception as ai_err:
                print(f"[Parser] OpenAI AI parsing failed or offline ({ai_err}), falling back to smart local heuristic parser.")

        # 2. Smart local heuristic parser (works 100% offline)
        parsed_resume = parse_raw_text_to_resume(raw_text, filename)
        return parsed_resume

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to parse resume: {str(e)}")

@app.post("/api/ai/rewrite", response_model=AIRewriteResponse)
async def ai_rewrite_endpoint(req: AIRewriteRequest):
    try:
        return await run_dual_pass_ai_rewrite(req)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI rewrite failed: {str(e)}")

@app.post("/api/ats/score")
async def ats_score_endpoint(req: ATSScoreRequest):
    try:
        return calculate_backend_ats_score(req.resumeData, req.targetJobDescription or "")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ATS calculation failed: {str(e)}")

@app.post("/api/resume/merge", response_model=ResumeMergeResponse)
async def resume_merge_endpoint(req: ResumeMergeRequest):
    try:
        return await run_resume_merger(req)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Resume merge failed: {str(e)}")

@app.post("/api/ai/generate-jd", response_model=GenerateJDResponse)
async def generate_jd_endpoint(req: GenerateJDRequest):
    try:
        return await generate_job_description_from_role(req.rolePrompt, req.apiKey)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"JD generation failed: {str(e)}")

@app.post("/api/export/docx")
async def export_docx_endpoint(resume: ResumeData):
    try:
        doc = Document()

        # Set 0.5in margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Header / Personal Info
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_name = p_name.add_run(resume.personalInfo.fullName.upper())
        run_name.font.name = "Arial"
        run_name.font.size = Pt(16)
        run_name.font.bold = True

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(resume.personalInfo.jobTitle)
        run_title.font.name = "Arial"
        run_title.font.size = Pt(11)

        # Contact line
        contacts = [c for c in [
            resume.personalInfo.phone,
            resume.personalInfo.email,
            resume.personalInfo.location,
            resume.personalInfo.linkedin
        ] if c]
        if contacts:
            p_contact = doc.add_paragraph()
            p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_contact = p_contact.add_run(" | ".join(contacts))
            run_contact.font.name = "Arial"
            run_contact.font.size = Pt(9.5)
            run_contact.font.color.rgb = RGBColor(100, 116, 139)

        # Sections
        for sec in resume.sections:
            if not sec.visible:
                continue

            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run_h = h.add_run(sec.title.upper())
            run_h.font.name = "Arial"
            run_h.font.size = Pt(11)
            run_h.font.bold = True
            run_h.font.color.rgb = RGBColor(30, 41, 59)

            # Strip HTML tags into paragraphs & bullets
            raw_content = sec.contentHtml
            # Check for bullets
            bullets = re.findall(r'<li>(.*?)</li>', raw_content, re.DOTALL)
            if bullets:
                for b in bullets:
                    clean_b = re.sub(r'<[^>]*>', '', b).strip()
                    p_b = doc.add_paragraph(clean_b, style='List Bullet')
                    p_b.paragraph_format.space_after = Pt(1)
            else:
                paragraphs = re.findall(r'<p>(.*?)</p>', raw_content, re.DOTALL)
                if paragraphs:
                    for p in paragraphs:
                        clean_p = re.sub(r'<[^>]*>', '', p).strip()
                        doc.add_paragraph(clean_p)
                else:
                    clean_all = re.sub(r'<[^>]*>', '', raw_content).strip()
                    doc.add_paragraph(clean_all)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        filename = f"{resume.personalInfo.fullName.replace(' ', '_')}_Resume.docx"
        return Response(
            content=file_stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Word export failed: {str(e)}")
